#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""make_deliverable_docx.py — 从 HTML 交付件派生 Word（.docx）版本。

单一事实来源：deliverables/DDPS_Deliverable.html（先由 make_deliverable.py 生成），
内容与 HTML 版逐字一致，不维护第二套文案。

动态部分的 Word 呈现策略（与打印版同构，全部内容平铺可见）：
- 页签（.tabs）   → 平铺为顺序小节：页签标签作为强调行「▍页签 · …」置于各面板前
- 逐帧走查（.wt） → 全部帧展开为连环画：帧标题 + 图 + 图注 + 说明文字（打印版同款）
- 折叠块（details）→ 直接展开（标题加粗 + 内容）
- 内联 SVG        → 无头 Edge 以页面同款 CSS 栅格化为 2x PNG 后嵌入（缓存 _parts/svg_png/）
- PNG 图（报告图）→ data URI 解码直接嵌入
- 目录            → Word 原生 TOC 域（打开后 F9 / 右键更新域生成）+ 标题样式驱动的导航窗格
- 「回目录」回链   → 不呈现（Word 有原生导航）

用法：
  python make_deliverable_docx.py --html deliverables/DDPS_Deliverable.html \
      --out deliverables/DDPS_Deliverable.docx
"""
import argparse
import base64
import hashlib
import io
import os
import re
import subprocess
import sys
import tempfile

from bs4 import BeautifulSoup, NavigableString, Tag

EDGE = r'C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe'
ACCENT = '0B63CE'
GRAY = '6B7A8D'
FILL_HEAD = 'EEF2F7'
FILL_CODE = 'F4F7FB'

# ---------------------------------------------------------------------------
# SVG 栅格化（无头 Edge，页面同款 CSS，2x，内容哈希缓存）
# ---------------------------------------------------------------------------

def rasterize_svgs(soup, style_css, cache_dir):
    """给文档里所有将使用的 <svg>（跳过 d-narrow）生成 PNG，返回 {svg元素id: png路径}。"""
    os.makedirs(cache_dir, exist_ok=True)
    mapping = {}
    for i, svg in enumerate(soup.find_all('svg')):
        cls = svg.get('class') or []
        if 'd-narrow' in cls:
            continue
        key = hashlib.md5(str(svg).encode('utf-8')).hexdigest()[:16]
        png = os.path.abspath(os.path.join(cache_dir, key + '.png'))
        if not os.path.exists(png):
            vb = [float(v) for v in svg.get('viewbox', svg.get('viewBox', '0 0 340 240')).split()[-2:]]
            W, H = int(round(vb[0])), int(round(vb[1]))
            wrapper = ('<!DOCTYPE html><html><head><meta charset="utf-8"><style>'
                       + style_css +
                       ' body{margin:0;background:#fff !important}'
                       ' svg{width:' + str(W) + 'px !important;height:' + str(H) +
                       'px !important;display:block !important}'
                       '</style></head><body>' + str(svg) + '</body></html>')
            with tempfile.TemporaryDirectory(prefix='svgwrap_') as td:
                hp = os.path.join(td, 'w.html')
                with open(hp, 'w', encoding='utf-8') as f:
                    f.write(wrapper)
                subprocess.run([EDGE, '--headless', '--disable-gpu', '--hide-scrollbars',
                                '--force-device-scale-factor=2',
                                '--window-size=%d,%d' % (W, H),
                                '--screenshot=%s' % png,
                                'file:///' + hp.replace('\\', '/')],
                               capture_output=True, timeout=120, check=False)
        if os.path.exists(png):
            mapping[svg] = png
        else:
            print('[docx] WARN: rasterize failed for svg #%d (viewBox=%s)' % (i, svg.get('viewbox', svg.get('viewBox'))))
    return mapping


# ---------------------------------------------------------------------------
# 文档构建
# ---------------------------------------------------------------------------

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def set_east_asia(style_or_run, name='微软雅黑'):
    rPr = style_or_run.element.get_or_add_rPr() if hasattr(style_or_run, 'element') else None
    if rPr is None:
        return
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), name)


def shade(el_pr, fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), fill)
    el_pr.append(shd)


class DocxBuilder:
    def __init__(self, svg_png):
        self.svg_png = svg_png
        self.doc = Document()
        self._setup()

    # ---- 样式与页面 ----
    def _setup(self):
        d = self.doc
        sec = d.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(1.9)
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)

        normal = d.styles['Normal']
        normal.font.name = 'Segoe UI'
        normal.font.size = Pt(10)
        set_east_asia(normal)
        normal.paragraph_format.space_after = Pt(4)
        normal.paragraph_format.line_spacing = 1.25

        for name, size, color, before, after in (
                ('Heading 1', 15.5, ACCENT, 16, 8),
                ('Heading 2', 13.0, ACCENT, 13, 6),
                ('Heading 3', 11.5, '12161C', 10, 5),
                ('Heading 4', 10.5, '3C4858', 8, 4)):
            st = d.styles[name]
            st.font.name = 'Segoe UI'
            st.font.size = Pt(size)
            st.font.bold = True
            st.font.color.rgb = RGBColor.from_string(color)
            set_east_asia(st)
            st.paragraph_format.space_before = Pt(before)
            st.paragraph_format.space_after = Pt(after)
            st.paragraph_format.keep_with_next = True

        for name in ('List Bullet', 'List Number', 'List Bullet 2', 'List Number 2'):
            try:
                st = d.styles[name]
                st.font.name = 'Segoe UI'
                st.font.size = Pt(10)
                set_east_asia(st)
            except KeyError:
                pass

    # ---- 基础构件 ----
    def para(self, style=None):
        return self.doc.add_paragraph(style=style)

    def add_runs(self, p, node, bold=False, italic=False, mono=False, gray=False):
        """把内联节点（文本/strong/code/sub/sup/a/br）展开成 Word run。"""
        for child in node.children:
            if isinstance(child, NavigableString):
                text = str(child)
                if text.strip() or text == ' ':
                    r = p.add_run(text)
                    r.bold = bold or None
                    r.italic = italic or None
                    if mono:
                        r.font.name = 'Consolas'
                        set_east_asia(r)
                    if gray:
                        r.font.color.rgb = RGBColor.from_string(GRAY)
            elif isinstance(child, Tag):
                cn = child.name
                ccls = child.get('class') or []
                if cn in ('strong', 'b'):
                    self.add_runs(p, child, bold=True)
                elif cn in ('em', 'i'):
                    self.add_runs(p, child, italic=True)
                elif cn in ('code',) or 'mono' in ccls:
                    self.add_runs(p, child, mono=True)
                elif cn == 'sub':
                    r = p.add_run(child.get_text())
                    r.font.subscript = True
                elif cn == 'sup':
                    r = p.add_run(child.get_text())
                    r.font.superscript = True
                elif cn == 'br':
                    p.add_run().add_break()
                elif cn == 'a':
                    self.add_runs(p, child)
                elif cn == 'span':
                    if 'sh' in ccls:        # 「可左右滑动」屏滚提示，Word 无意义
                        continue
                    self.add_runs(p, child, gray=('mut' in ccls) or gray,
                                  mono=('mono' in ccls) or mono)
                else:
                    self.add_runs(p, child)

    def text_of(self, el):
        return re.sub(r'\s+', ' ', el.get_text()).strip()

    def heading(self, el, level):
        text = self.text_of(el)
        text = re.sub(r'^(\d+(?:\.\d+)*)(?=\S)', r'\1 ', text)   # “6.5对比…” → “6.5 对比…”
        h = self.doc.add_heading(level=level)
        h.add_run(text)

    def caption(self, text, center=False):
        p = self.para()
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor.from_string(GRAY)
        return p

    def leadin(self, text, color=ACCENT):
        """强调行（页签标签 / 帧标题）。"""
        p = self.para()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor.from_string(color)
        return p

    def add_svg(self, svg, width_cm):
        png = self.svg_png.get(svg)
        if png:
            p = self.para()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(png, width=Cm(width_cm))

    def add_img(self, img, width_cm=16.5):
        src = img.get('src') or ''
        m = re.match(r'data:image/png;base64,(.*)', src, re.S)
        if not m:
            return
        data = base64.b64decode(m.group(1))
        p = self.para()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(io.BytesIO(data), width=Cm(width_cm))

    def add_table(self, tbl):
        rows = tbl.find_all('tr', recursive=False)
        if not rows:
            return
        cap = tbl.find('caption')
        if cap:
            t = self.text_of(cap)
            t = re.sub(r'\s*·\s*可左右滑动\s*$', '', t)
            self.caption('表 · ' + t)
        header_cells = rows[0].find_all(['th', 'td'], recursive=False)
        ncol = max(len(r.find_all(['th', 'td'], recursive=False)) for r in rows)
        wt = self.doc.add_table(rows=len(rows), cols=ncol)
        wt.style = 'Table Grid'
        wt.alignment = WD_TABLE_ALIGNMENT.CENTER
        wt.autofit = True
        for ri, row in enumerate(rows):
            cells = row.find_all(['th', 'td'], recursive=False)
            for ci in range(ncol):
                cell = wt.cell(ri, ci)
                el = cells[ci] if ci < len(cells) else None
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(1)
                is_head = ri == 0 or (el is not None and el.name == 'th')
                if el is not None:
                    self.add_runs(p, el, bold=is_head)
                    for r in p.runs:
                        r.font.size = Pt(8.2)
                if is_head:
                    shade(cell._tc.get_or_add_tcPr(), FILL_HEAD)
                if el is not None and 'n' in (el.get('class') or []):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # 隔行浅灰
        for ri in range(2, len(rows), 2):
            for ci in range(ncol):
                shade(wt.cell(ri, ci)._tc.get_or_add_tcPr(), 'F8FAFD')

    def add_pre(self, pre):
        for line in pre.get_text().split('\n'):
            p = self.para()
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(line if line.strip() else ' ')
            r.font.name = 'Consolas'
            r.font.size = Pt(8.2)
            set_east_asia(r)
            shade(p._p.get_or_add_pPr(), FILL_CODE)

    def add_kpis(self, div):
        cards = div.find_all('div', class_='kpi', recursive=False)
        if not cards:
            return
        wt = self.doc.add_table(rows=2, cols=len(cards))
        wt.style = 'Table Grid'
        wt.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, c in enumerate(cards):
            v, l = c.find('div', class_='v'), c.find('div', class_='l')
            pv = wt.cell(0, i).paragraphs[0]
            rv = pv.add_run(self.text_of(v) if v else '')
            rv.bold = True
            rv.font.size = Pt(13)
            rv.font.color.rgb = RGBColor.from_string(ACCENT)
            pv.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pl = wt.cell(1, i).paragraphs[0]
            rl = pl.add_run(self.text_of(l) if l else '')
            rl.font.size = Pt(8.0)
            rl.font.color.rgb = RGBColor.from_string(GRAY)

    def _toc_page(self):
        """标题/引言之后的目录页：Word 原生 TOC 域（H1-H2），打开后 F9 更新。"""
        self.doc.add_page_break()
        p = self.para()
        r = p.add_run('目录')
        r.bold = True
        r.font.size = Pt(14)
        r.font.color.rgb = RGBColor.from_string(ACCENT)
        pf = self.para()
        fld = OxmlElement('w:fldSimple')
        fld.set(qn('w:instr'), r' TOC \o "1-2" \h \z \u ')
        rr = OxmlElement('w:r')
        tt = OxmlElement('w:t')
        tt.text = '（目录域：打开文档后按 F9 或右键「更新域」生成章节目录与页码）'
        rr.append(tt)
        fld.append(rr)
        pf._p.append(fld)
        self.doc.add_page_break()

    # ---- 主遍历 ----
    SKIP_CLASSES = ('back-toc', 'wt-bar', 'd-narrow', 'tab-bar', 'toc')

    def skip(self, el):
        if not isinstance(el, Tag):
            return True
        if el.name in ('script', 'style', 'button', 'nav'):
            return True
        cls = set(el.get('class') or [])
        return bool(cls & set(self.SKIP_CLASSES))

    def emit_children(self, el):
        for c in el.children:
            self.emit(c)

    def emit(self, el):
        if self.skip(el):
            return
        name = el.name
        cls = el.get('class') or []

        if name == 'header':
            self.emit_children(el)
            self._toc_page()
        elif name == 'h1':
            t = self.text_of(el)
            p = self.para()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(t)
            r.bold = True
            r.font.size = Pt(19)
            r.font.color.rgb = RGBColor.from_string(ACCENT)
        elif name == 'div' and 'eyebrow' in cls:
            p = self.para()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(self.text_of(el))
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor.from_string(GRAY)
        elif name == 'div' and 'scope' in cls:
            p = self.para()
            self.add_runs(p, el)
        elif name in ('h2', 'h3', 'h4'):
            self.heading(el, {'h2': 1, 'h3': 2, 'h4': 3}[name])
        elif name == 'h5':
            p = self.para()
            r = p.add_run(self.text_of(el))
            r.bold = True
            r.font.color.rgb = RGBColor.from_string(ACCENT)
        elif name == 'p':
            if not self.text_of(el):
                return
            p = self.para()
            self.add_runs(p, el)
        elif name in ('ul', 'ol'):
            self.emit_list(el, name, 1)
        elif name == 'table':
            self.add_table(el)
        elif name == 'pre':
            self.add_pre(el)
        elif name == 'figure':
            self.emit_figure(el)
        elif name == 'details':
            s = el.find('summary')
            if s:
                p = self.para()
                r = p.add_run('▸ ' + self.text_of(s))
                r.bold = True
                r.font.color.rgb = RGBColor.from_string('3C4858')
                p.paragraph_format.space_before = Pt(8)
            body = el.find('div', class_='fold-body')
            if body:
                self.emit_children(body)
        elif name == 'div' and 'tabs' in cls:
            self.emit_tabs(el)
        elif name == 'div' and 'wt' in cls and 'wt-frame' not in cls:
            self.emit_walkthrough(el)
        elif name == 'div' and 'kpis' in cls:
            self.add_kpis(el)
        elif name == 'div' and 'wt-frame' in cls and 'on' in cls:
            # 独立复用走查栅格布局的块（如 §3.1 训练数据卡）：两栏顺序平铺
            self.emit_children(el)
        elif name == 'img':
            if not self._is_narrow(el):
                self.add_img(el)
        elif name == 'svg':
            if not self._is_narrow(el):
                self.add_svg(el, 16.0)
        elif name == 'footer':
            self.emit_children(el)
        elif isinstance(el, Tag):
            self.emit_children(el)

    def emit_list(self, el, name, depth):
        style = ('List Bullet' if name == 'ul' else 'List Number') + (' 2' if depth > 1 else '')
        try:
            self.doc.styles[style]
        except KeyError:
            style = None
        for li in el.find_all('li', recursive=False):
            p = self.para(style=style)
            for sub in li.children:                      # 列表项内联内容（跳过嵌套列表）
                if isinstance(sub, Tag) and sub.name in ('ul', 'ol'):
                    continue
                if isinstance(sub, NavigableString):
                    t = str(sub)
                    if t.strip():
                        p.add_run(t)
                else:
                    self.add_runs(p, sub)
            for sub in li.find_all(['ul', 'ol'], recursive=False):
                self.emit_list(sub, sub.name, depth + 1)

    def _is_narrow(self, el):
        """d-narrow 响应式变体（屏窄版）——Word 版只取 d-wide。"""
        p = el
        while p is not None and getattr(p, 'name', None):
            if 'd-narrow' in (p.get('class') or []):
                return True
            p = p.parent
        return False

    def emit_figure(self, fig):
        title = fig.find('div', class_='fig-title')
        if title:
            p = self.para()
            r = p.add_run(self.text_of(title))
            r.bold = True
            p.paragraph_format.keep_with_next = True
        for img in fig.find_all('img'):
            if not self._is_narrow(img):
                self.add_img(img, 16.5)
        for svg in fig.find_all('svg'):
            if not self._is_narrow(svg):
                self.add_svg(svg, 16.0)
        cap = fig.find('figcaption')
        if cap:
            self.caption(self.text_of(cap), center=True)

    def emit_tabs(self, wrap):
        """页签 → 顺序小节：每个面板前放「▍页签 · 标签」强调行。"""
        labels = [self.text_of(b) for b in wrap.find_all('button', class_='tab-btn')]
        panels = wrap.find_all('div', class_='tab-panel', recursive=False)
        for i, panel in enumerate(panels):
            if i < len(labels):
                self.leadin('▍页签 %d/%d · %s' % (i + 1, len(panels), labels[i]))
            self.emit_children(panel)

    def emit_walkthrough(self, wrap):
        """逐帧走查 → 连环画：帧标题 + 图 + 图注 + 说明文字（打印版同款）。"""
        frames = wrap.find_all('div', class_='wt-frame', recursive=False)
        for f in frames:
            self.leadin('▍' + (f.get('data-title') or ''))
            halves = f.find_all('div', recursive=False)
            for half in halves:
                if 'wt-right' in (half.get('class') or []):
                    for c in half.children:
                        if isinstance(c, Tag) and c.name == 'h5':
                            p = self.para()
                            r = p.add_run(self.text_of(c))
                            r.bold = True
                            r.font.color.rgb = RGBColor.from_string(ACCENT)
                        elif isinstance(c, Tag):
                            self.emit(c)
                else:
                    for svg in half.find_all('svg'):
                        self.add_svg(svg, 14.2)
                    cap = half.find('p', class_='wt-cap')
                    if cap:
                        self.caption(self.text_of(cap))


# ---------------------------------------------------------------------------
# 入口
# ---------------------------------------------------------------------------

def build(html_path, out_path):
    html = open(html_path, encoding='utf-8').read()
    soup = BeautifulSoup(html, 'html.parser')
    style_css = re.search(r'<style>(.*?)</style>', html, re.S).group(1)

    cache = os.path.join('_parts', 'svg_png')
    print('[docx] rasterizing inline SVGs (headless Edge, 2x) ...')
    svg_png = rasterize_svgs(soup, style_css, cache)
    print('[docx] %d svg -> png' % len(svg_png))

    b = DocxBuilder(svg_png)
    for el in soup.body.children:
        b.emit(el)

    b.doc.core_properties.title = 'DDPS 交付说明 — 收发端均衡代理寻优（Word 版）'
    b.doc.core_properties.comments = '由 deliverables/DDPS_Deliverable.html 派生：make_deliverable_docx.py'
    b.doc.save(out_path)
    print('[docx] written %s (%.1f MB)' % (out_path, os.path.getsize(out_path) / 1048576))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--html', default='deliverables/DDPS_Deliverable.html')
    ap.add_argument('--out', default='deliverables/DDPS_Deliverable.docx')
    a = ap.parse_args()
    build(a.html, a.out)


if __name__ == '__main__':
    main()
